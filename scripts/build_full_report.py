import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_ORIENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(27, 94, 32) # Dark agricultural green
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(46, 125, 50)
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(33, 33, 33)
    return p

def add_body_p(doc, text, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(33, 33, 33)
    return p

def add_code_block(doc, title, code_str):
    p_t = doc.add_paragraph()
    p_t.paragraph_format.space_before = Pt(6)
    p_t.paragraph_format.space_after = Pt(2)
    p_t.paragraph_format.keep_with_next = True
    run_t = p_t.add_run(title)
    run_t.font.name = 'Calibri'
    run_t.font.size = Pt(9)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(46, 125, 50)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F5F5F5")
    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="2E7D32"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_str.strip())
    run.font.name = 'Consolas'
    run.font.size = Pt(7.5)
    run.font.color.rgb = RGBColor(38, 50, 56)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_figure(doc, img_path, title, live_link=None, qr_path=None, width=Inches(5.6)):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.paragraph_format.keep_with_next = True
        p_img.add_run().add_picture(img_path, width=width)
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(2)
        p_cap.paragraph_format.keep_with_next = True
        run_cap = p_cap.add_run(title)
        run_cap.font.name = 'Calibri'
        run_cap.font.size = Pt(8.5)
        run_cap.font.bold = True
        run_cap.font.color.rgb = RGBColor(50, 50, 50)
        
        if live_link or qr_path:
            p_sub = doc.add_paragraph()
            p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sub.paragraph_format.space_after = Pt(6)
            if live_link:
                run_l = p_sub.add_run(f"🔗 Interactive Live Map: {live_link}  ")
                run_l.font.name = 'Calibri'
                run_l.font.size = Pt(7.5)
                run_l.font.italic = True
                run_l.font.color.rgb = RGBColor(21, 101, 192)
            if qr_path and os.path.exists(qr_path):
                p_sub.add_run().add_picture(qr_path, width=Inches(0.6))
    else:
        p = doc.add_paragraph(f"[Image Missing: {img_path}]")
        p.font.color.rgb = RGBColor(200, 0, 0)

print("Report generator setup initialized.")
