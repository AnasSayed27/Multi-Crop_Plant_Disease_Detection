import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_ORIENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def add_callout_box(doc, title, text_lines, bg_color="F1F8E9", border_color="2E7D32"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    # Left border only
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run_t = p.add_run(title + "\n")
    run_t.font.name = 'Calibri'
    run_t.font.size = Pt(9.5)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(46, 125, 50)
    
    for l in text_lines:
        run_l = p.add_run(l + "\n")
        run_l.font.name = 'Calibri'
        run_l.font.size = Pt(8.5)
        run_l.font.color.rgb = RGBColor(50, 50, 50)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code_block(doc, snippet_title, code_text):
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(2)
    run_t = p_title.add_run(snippet_title)
    run_t.font.name = 'Calibri'
    run_t.font.size = Pt(9.5)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(33, 33, 33)
    
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F5F5F5")
    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:left w:val="single" w:sz="18" w:space="0" w:color="2E7D32"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_text.strip())
    run.font.name = 'Consolas'
    run.font.size = Pt(7.5)
    run.font.color.rgb = RGBColor(38, 50, 56)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_figure_with_qr(doc, img_path, fig_title, qr_path=None, live_link=None, width=Inches(5.8)):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture(img_path, width=width)
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(4)
        run_cap = p_cap.add_run(fig_title)
        run_cap.font.name = 'Calibri'
        run_cap.font.size = Pt(8.5)
        run_cap.font.bold = True
        run_cap.font.color.rgb = RGBColor(60, 60, 60)
        
        if live_link:
            p_link = doc.add_paragraph()
            p_link.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_link.paragraph_format.space_after = Pt(8)
            run_link = p_link.add_run(f"🔗 [Interactive Live Diagram URL]: {live_link}")
            run_link.font.name = 'Calibri'
            run_link.font.size = Pt(7.5)
            run_link.font.italic = True
            run_link.font.color.rgb = RGBColor(21, 101, 192)
    else:
        p_err = doc.add_paragraph(f"[Figure missing: {img_path}]")
        p_err.font.color.rgb = RGBColor(200, 0, 0)

print("Helper functions defined successfully.")
