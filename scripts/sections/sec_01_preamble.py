# scripts/sections/sec_01_preamble.py
import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_ORIENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def init_document():
    doc = docx.Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.left_margin = Inches(1.25)   # 1.25 in binding margin
    sec.right_margin = Inches(1.0)
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)

    # Configure base style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(12)
    style_normal.font.color.rgb = RGBColor(38, 50, 56) # charcoal
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(4.5)
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Header & Footer setup
    header = sec.header
    p_hdr = header.paragraphs[0]
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_hdr = p_hdr.add_run("Multi-Crop Plant Disease Detection & Advisory System (DPD ViT-Base)")
    r_hdr.font.name = 'Calibri'
    r_hdr.font.size = Pt(8.5)
    r_hdr.font.italic = True
    r_hdr.font.color.rgb = RGBColor(120, 144, 156)

    footer = sec.footer
    p_ftr = footer.paragraphs[0]
    p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ftr = p_ftr.add_run("Department of Computer Science | B.Sc. CS Dissertation")
    r_ftr.font.name = 'Calibri'
    r_ftr.font.size = Pt(8.5)
    r_ftr.font.color.rgb = RGBColor(120, 144, 156)

    # Helper closures bound to doc
    def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f"""
            <w:tcMar {nsdecls('w')}>
                <w:top w:w="{top}" w:type="dxa"/>
                <w:bottom w:w="{bottom}" w:type="dxa"/>
                <w:left w:w="{left}" w:type="dxa"/>
                <w:right w:w="{right}" w:type="dxa"/>
            </w:tcMar>
        """)
        tcPr.append(tcMar)

    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_row_cant_split(row):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    def set_table_header(row):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = RGBColor(27, 94, 32) # Dark Green
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(46, 125, 50) # Forest Green
        return p

    def add_heading_3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(33, 33, 33) # Charcoal
        return p

    def add_body_p(text, space_after=4.5):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pPr = p._p.get_or_add_pPr()
        pPr.append(parse_xml(f'<w:widowControl {nsdecls("w")}/>'))
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(38, 50, 56)
        return p

    def add_bullet_p(text, bold_prefix="", space_after=3):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pPr = p._p.get_or_add_pPr()
        pPr.append(parse_xml(f'<w:widowControl {nsdecls("w")}/>'))
        if bold_prefix:
            r_b = p.add_run(bold_prefix)
            r_b.font.name = 'Calibri'
            r_b.font.size = Pt(12)
            r_b.font.bold = True
            r_b.font.color.rgb = RGBColor(46, 125, 50)
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(38, 50, 56)
        return p

    def add_code_block(title, code_str):
        p_t = doc.add_paragraph()
        p_t.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_t.paragraph_format.space_before = Pt(8)
        p_t.paragraph_format.space_after = Pt(3)
        p_t.paragraph_format.keep_with_next = True
        run_t = p_t.add_run(title)
        run_t.font.name = 'Calibri'
        run_t.font.size = Pt(10)
        run_t.font.bold = True
        run_t.font.color.rgb = RGBColor(46, 125, 50)

        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.0)
        set_cell_background(cell, "F8F9FA")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        set_row_cant_split(tbl.rows[0])

        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(f"""
            <w:tcBorders {nsdecls('w')}>
                <w:top w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="2E7D32"/>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>
                <w:right w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>
            </w:tcBorders>
        """)
        tcPr.append(tcBorders)

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(code_str.strip())
        run.font.name = 'Consolas'
        run.font.size = Pt(9.0)
        run.font.color.rgb = RGBColor(38, 50, 56)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_callout_box(title, text, bg_color="F1F8E9", border_color="2E7D32"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.0)
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        set_row_cant_split(tbl.rows[0])

        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(f"""
            <w:tcBorders {nsdecls('w')}>
                <w:top w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
                <w:right w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
            </w:tcBorders>
        """)
        tcPr.append(tcBorders)

        # Title paragraph inside box (strictly LEFT-aligned, never justified)
        p_title = cell.paragraphs[0]
        p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_title.paragraph_format.space_before = Pt(2)
        p_title.paragraph_format.space_after = Pt(3)
        p_title.paragraph_format.line_spacing = 1.15
        r_t = p_title.add_run(title.strip())
        r_t.font.name = 'Calibri'
        r_t.font.size = Pt(10.5)
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(27, 94, 32)

        # Body paragraph inside box (strictly LEFT-aligned)
        p_body = cell.add_paragraph()
        p_body.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_body.paragraph_format.space_before = Pt(0)
        p_body.paragraph_format.space_after = Pt(2)
        p_body.paragraph_format.line_spacing = 1.15

        # Code / Logic detection
        is_code_like = any(line.strip().startswith(('IF ', 'ELSE:', 'is_', 'status =', 'crop =', 'disease =', 'advisory =', 'Category:')) for line in text.split('\n'))
        r_b = p_body.add_run(text.strip())
        r_b.font.name = 'Consolas' if is_code_like else 'Calibri'
        r_b.font.size = Pt(9.0) if is_code_like else Pt(10)
        r_b.font.color.rgb = RGBColor(38, 50, 56)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_figure(img_path, caption, live_link=None, qr_path=None, width=Inches(5.2), page_break_before=False):
        if page_break_before:
            doc.add_page_break()
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(2)
            p_img.paragraph_format.keep_with_next = True
            p_img.add_run().add_picture(img_path, width=width)

            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(2)
            p_cap.paragraph_format.keep_with_next = True if (live_link or qr_path) else False
            run_cap = p_cap.add_run(caption)
            run_cap.font.name = 'Calibri'
            run_cap.font.size = Pt(9.5)
            run_cap.font.bold = True
            run_cap.font.color.rgb = RGBColor(50, 50, 50)

            if live_link or qr_path:
                p_sub = doc.add_paragraph()
                p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_sub.paragraph_format.space_after = Pt(6)
                if live_link:
                    run_l = p_sub.add_run(f"🔗 [Interactive Live Diagram]: {live_link}  ")
                    run_l.font.name = 'Calibri'
                    run_l.font.size = Pt(8.5)
                    run_l.font.italic = True
                    run_l.font.color.rgb = RGBColor(21, 101, 192)
                if qr_path and os.path.exists(qr_path):
                    p_sub.add_run().add_picture(qr_path, width=Inches(0.6))
        else:
            p = doc.add_paragraph(f"[Image Missing: {img_path}]")
            p.font.color.rgb = RGBColor(200, 0, 0)

    def add_table(headers, data, col_widths=None, caption=None, page_break_before=False):
        if page_break_before:
            doc.add_page_break()
        if caption:
            p_cap = doc.add_paragraph()
            p_cap.paragraph_format.space_before = Pt(8)
            p_cap.paragraph_format.space_after = Pt(3)
            p_cap.paragraph_format.keep_with_next = True
            r_c = p_cap.add_run(caption)
            r_c.font.name = 'Calibri'
            r_c.font.size = Pt(10)
            r_c.font.bold = True
            r_c.font.color.rgb = RGBColor(27, 94, 32)

        tbl = doc.add_table(rows=len(data)+1, cols=len(headers))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Header Row
        hdr_row = tbl.rows[0]
        set_table_header(hdr_row)
        set_row_cant_split(hdr_row)
        for idx, h_text in enumerate(headers):
            cell = hdr_row.cells[idx]
            if col_widths and idx < len(col_widths):
                cell.width = col_widths[idx]
            cell.text = h_text
            set_cell_background(cell, "2E7D32")
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            if p.runs:
                p.runs[0].font.name = 'Calibri'
                p.runs[0].font.size = Pt(9.5)
                p.runs[0].font.bold = True
                p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

        # Data Rows
        for r_idx, row_vals in enumerate(data):
            row = tbl.rows[r_idx+1]
            set_row_cant_split(row)
            bg_color = "FFFFFF" if r_idx % 2 == 0 else "F9FBE7"
            for c_idx, val in enumerate(row_vals):
                cell = row.cells[c_idx]
                if col_widths and c_idx < len(col_widths):
                    cell.width = col_widths[c_idx]
                cell.text = str(val)
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=50, bottom=50, left=90, right=90)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                if p.runs:
                    p.runs[0].font.name = 'Calibri'
                    p.runs[0].font.size = Pt(9.0)
                    p.runs[0].font.color.rgb = RGBColor(38, 50, 56)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        return tbl

    helpers = {
        'doc': doc,
        'add_heading_1': add_heading_1,
        'add_heading_2': add_heading_2,
        'add_heading_3': add_heading_3,
        'add_body_p': add_body_p,
        'add_bullet_p': add_bullet_p,
        'add_code_block': add_code_block,
        'add_callout_box': add_callout_box,
        'add_figure': add_figure,
        'add_table': add_table,
        'set_cell_background': set_cell_background,
        'set_cell_margins': set_cell_margins,
        'set_row_cant_split': set_row_cant_split,
        'set_table_header': set_table_header
    }
    return doc, helpers