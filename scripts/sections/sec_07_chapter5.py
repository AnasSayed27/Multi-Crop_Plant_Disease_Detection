# scripts/sections/sec_07_chapter5.py
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_chapter5(doc, helpers):
    add_h1 = helpers['add_heading_1']
    add_h2 = helpers['add_heading_2']
    add_h3 = helpers['add_heading_3']
    add_p = helpers['add_body_p']
    add_bullet = helpers['add_bullet_p']
    add_code = helpers['add_code_block']
    add_callout = helpers['add_callout_box']

    # -------------------------------------------------------------
    # CHAPTER 5: IMPLEMENTATION AND METHODOLOGIES
    # -------------------------------------------------------------
    add_h1("CHAPTER 5: IMPLEMENTATION AND METHODOLOGIES")

    add_h2("5.1 Project Directory Structure & Modular Decomposition")
    add_p("The software codebase is organized according to enterprise software engineering standards, enforcing a strict separation of concerns between deep learning modeling, web service gateways, relational persistence, clinical reporting, and static front-end assets. The top-level project directory structure is illustrated below:")

    dir_structure = """Potato_disease/
├── app.py                      # Asynchronous FastAPI web application and REST routing
├── dpd_model.py                # PyTorch DPD ViT-Base inference engine and candidate ranking
├── database.py                 # SQLite WAL connection management, hashing, and queries
├── pdf_generator.py            # ReportLab clinical diagnostic PDF flowable compiler
├── static/                     # Front-end user interface assets
│   ├── css/style.css           # Custom responsive styles and animations
│   ├── js/main.js              # Client-side drag-and-drop, fetch API, and DOM binding
│   └── index.html              # Responsive farmer dashboard and diagnostic modal
├── models_assets/              # Serialized neural weights and domain taxonomy schemas
│   ├── model_b_partial_adapted.pth  # Fine-tuned DPD ViT-Base PyTorch state dictionary (343 MB)
│   ├── dpd_55_plants.json      # Botanical index mapping for 55 plant species
│   ├── dpd_175_diseases.json   # Pathological index mapping for 175 disease conditions
│   ├── dpd_crop_rankings.json  # Hierarchical taxonomy of 333 valid crop-disease pairs
│   └── disease_info.json       # Clinical 5-pillar agronomic treatment knowledge base
├── tests/                      # Automated test suite (PyTest framework)
│   ├── test_api.py             # End-to-end REST endpoint and security contract tests
│   ├── test_dpd_model.py       # Tensor shape, confidence calibration, and OOD tests
│   └── test_taxonomy_contracts.py  # Validation of 333-pair completeness and advisory data
└── docs/report_figures/        # High-resolution architectural and UML diagram assets"""
    add_code("Directory Hierarchy: Project File Organization", dir_structure)

    add_h2("5.2 Deep Learning Inference Pipeline (dpd_model.py)")
    add_p("The core neural inference pipeline is implemented in dpd_model.py. The module defines the DPDViTDualHead PyTorch architecture, wraps it in a singleton DPDInferenceEngine, loads the pre-trained weights from model_b_partial_adapted.pth, and manages hardware execution across CUDA and CPU backends.")

    add_h3("5.2.1 DPDViTDualHead PyTorch Module")
    add_p("The dual-head architecture utilizes a pre-trained vit_base_patch16_224 backbone sourced from the timm library. The original 1,000-class ImageNet classification head is stripped, and two independent linear projections are attached to the 768-dimensional output of the [CLS] token:")

    snippet_5_1 = """# File: dpd_model.py (Lines 142-155)
# Architecture: DPD ViT-Base Dual-Head PyTorch Module
class DPDViTDualHead(nn.Module):
    def __init__(self, base_model, num_plants=55, num_diseases=175):
        super().__init__()
        self.base_model = base_model
        # Decoupled projection heads sharing the 768-dim ViT representation
        self.linear_plant = nn.Linear(768, num_plants)
        self.linear_disease = nn.Linear(768, num_diseases)

    def forward(self, x: torch.Tensor):
        features = self.base_model(x)  # 768-dimensional latent vector
        out_plant = self.linear_plant(features)      # Logits across 55 crops
        out_disease = self.linear_disease(features)  # Logits across 175 diseases
        return out_plant, out_disease"""
    add_code("Snippet 5.1: DPDViTDualHead PyTorch Architecture Definition", snippet_5_1)

    add_h3("5.2.2 Tensor Preprocessing & Normalization")
    add_p("Input foliar images of arbitrary dimensions and aspect ratios are transformed into model-ready tensors via a standardized torchvision transform pipeline. Images are resized to 224x224 pixels using bicubic interpolation, converted to single-precision float tensors in [0.0, 1.0], and normalized using standard ImageNet channel means and standard deviations:")

    snippet_5_2 = """# File: dpd_model.py (Lines 160-170)
# Preprocessing: Standardized 224x224 ImageNet Transform Pipeline
self.transform = transforms.Compose([
    transforms.Resize((224, 224), interpolation=transforms.InterpolationMode.BICUBIC),
    transforms.ToTensor(),
    transforms.Normalize(
        mean=[0.485, 0.456, 0.406],
        std=[0.229, 0.224, 0.225]
    )
])"""
    add_code("Snippet 5.2: Image Normalization & Tensor Transformation Pipeline", snippet_5_2)

    add_h3("5.2.3 Formal Dual-Head Inference & Confidence Calibration Algorithm")
    add_p("Algorithm 1 details the step-by-step computational logic executed during foliar diagnosis, outlining tensor normalization, forward pass execution, candidate joint likelihood calculation across the 333 taxonomy matrix, and Out-of-Distribution gating:")

    algo_1_text = """ALGORITHM 1: DPD ViT-Base Dual-Head Inference & Confidence Calibration
Input:  Foliar image x, Taxonomy matrix Omega (333 pairs), Gating threshold theta = 40.0%
Output: Diagnostic result dictionary D with Top-1, Top-3 differential candidates, and Advisory

1:  tensor_x <- Normalize(Resize(x, 224, 224))
2:  z <- ViT_Backbone(tensor_x)                             // Latent vector in R^768
3:  logits_plant <- Linear_Plant(z)                          // Logits in R^55
4:  logits_disease <- Linear_Disease(z)                      // Logits in R^175
5:  prob_plant <- Softmax(logits_plant)
6:  prob_disease <- Softmax(logits_disease)
7:  candidates <- []
8:  FOR EACH (plant_idx, disease_idx, pair_slug) IN Omega DO:
9:      P_P <- prob_plant[plant_idx]
10:     P_D <- prob_disease[disease_idx]
11:     joint_prob <- P_P * P_D
12:     calibrated_conf <- sqrt(max(0, joint_prob)) * 100.0
13:     candidates.append({slug: pair_slug, conf: calibrated_conf, P_P: P_P, P_D: P_D})
14: END FOR
15: candidates <- SortDescending(candidates, key=conf)
16: top1 <- candidates[0]
17: top3 <- candidates[0:3]
18: IF top1.conf < theta THEN:
19:     RETURN {is_background: True, status: 'Uncertain', crop: 'Non-Crop', advisory: null}
20: ELSE:
21:     advisory <- Match_5_Pillar_Advisory(top1.slug)
22:     RETURN {is_background: False, crop: top1.crop, disease: top1.disease, conf: top1.conf, top3: top3, advisory: advisory}
23: END IF"""
    add_code("Algorithm 1: Dual-Head Forward Pass & Geometric Calibration Pseudocode", algo_1_text)

    snippet_5_3 = """# File: dpd_model.py (Lines 235-255)
# Calibration: Joint Likelihood Evaluation & Monotonic Candidate Sorting
candidates = []
for p in self.supported_pairs:
    prob_p = float(prob_plant[p['plant_idx']])
    prob_d = float(prob_disease[p['disease_idx']])
    joint_prob = prob_p * prob_d
    # Calibrated geometric mean confidence percentage
    calibrated_conf = float(np.sqrt(max(0.0, joint_prob)) * 100.0)
    candidates.append({
        'crop': p['crop'],
        'disease': p['disease'],
        'raw_class': p['raw_class'],
        'confidence': calibrated_conf,
        'plant_prob': prob_p,
        'disease_prob': prob_d
    })

# Strict descending sort guarantees monotonic candidate ordering
candidates.sort(key=lambda x: x['confidence'], reverse=True)
top1 = candidates[0]
top3 = candidates[:3]"""
    add_code("Snippet 5.3: Joint Likelihood & Monotonic Candidate Sorting Implementation", snippet_5_3)

    add_h2("5.3 Asynchronous Web Server & REST API Implementation (app.py)")
    add_p("The web service gateway is engineered using FastAPI, leveraging asynchronous Python coroutines (async/await) and the Uvicorn ASGI server to sustain high-concurrency workloads without thread starvation.")

    add_h3("5.3.1 Lifespan Model Management")
    add_p("To avoid reloading the 343 MB PyTorch model checkpoint on every incoming HTTP request, FastAPI's @asynccontextmanager lifespan handler initializes the DPDInferenceEngine once during server bootstrap and injects it into app.state:")

    snippet_5_4 = """# File: app.py (Lines 45-62)
# Architecture: Single-Instance Lifespan Management
@asynccontextmanager
async def lifespan(app: FastAPI):
    # Startup: Initialize DB schema and load neural weights into memory
    print("[Bootstrap] Initializing SQLite WAL database...")
    init_db()
    print("[Bootstrap] Loading DPD ViT-Base Dual-Head model into GPU/CPU memory...")
    app.state.model = DPDInferenceEngine(
        model_path="models_assets/model_b_partial_adapted.pth",
        assets_dir="models_assets"
    )
    print("[Bootstrap] System online and ready for incoming diagnostic traffic.")
    yield
    # Shutdown: Clean up tensor memory
    print("[Shutdown] Releasing model resources.")"""
    add_code("Snippet 5.4: FastAPI Lifespan Handler & Single-Instance Engine Cache", snippet_5_4)

    add_h3("5.3.2 Asynchronous Foliar Diagnosis Endpoint (/predict)")
    add_p("The primary inference route /predict accepts multipart/form-data image streams, validates MIME type and file size, invokes the neural engine in a threadpool executor, records scan metadata to SQLite, and returns structured JSON:")

    snippet_5_5 = """# File: app.py (Lines 110-145)
# Controller: Asynchronous Foliar Diagnosis Endpoint
@app.post("/predict")
async def predict_foliar_disease(
    file: UploadFile = File(...),
    current_user: Optional[dict] = Depends(get_optional_current_user)
):
    # Security Boundary 1: MIME Type Verification
    if file.content_type not in ["image/jpeg", "image/png", "image/webp"]:
        raise HTTPException(status_code=400, detail="Invalid file type. Please upload JPEG/PNG/WebP.")

    # Security Boundary 2: 10 MB Payload Guard
    contents = await file.read()
    if len(contents) == 0:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")
    if len(contents) > 10 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="File exceeds maximum allowed upload size of 10 MB.")

    # Image Decoding & Threadpool Inference Execution
    image = Image.open(io.BytesIO(contents)).convert("RGB")
    result = await asyncio.to_thread(app.state.model.predict, image)

    # Persistence to SQLite WAL Database
    user_id = current_user["id"] if current_user else 1
    prediction_id = save_prediction(
        user_id=user_id,
        crop=result["crop"],
        disease=result["disease"],
        confidence=result["confidence"],
        image_path=saved_path
    )
    result["prediction_id"] = prediction_id
    return JSONResponse(content=result)"""
    add_code("Snippet 5.5: Asynchronous Foliar Diagnosis Controller with Security Guards", snippet_5_5)

    add_h2("5.4 Database Access Layer & Transaction Management (database.py)")
    add_p("The database access layer (database.py) configures SQLite3 for enterprise multi-reader concurrency using Write-Ahead Logging (WAL) mode. When WAL mode is active, read transactions query snapshot states without blocking concurrent write commits:")

    snippet_5_6 = """# File: database.py (Lines 89-98)
# Persistence: SQLite WAL Connection Factory with Timeout Guard
def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    # Enforce relational integrity and concurrency pragmas
    conn.execute("PRAGMA foreign_keys = ON;")
    conn.execute("PRAGMA journal_mode = WAL;")
    conn.execute("PRAGMA busy_timeout = 5000;")
    return conn"""
    add_code("Snippet 5.6: SQLite WAL Connection Factory and Concurrency PRAGMAs", snippet_5_6)

    snippet_5_7 = """# File: database.py (Lines 160-185)
# Persistence: Scan Persistence and Historical Analytics Query
def save_prediction(user_id: int, crop: str, disease: str, confidence: float, image_path: str) -> int:
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute(\"\"\"
        INSERT INTO predictions (user_id, crop, disease, confidence, image_path)
        VALUES (?, ?, ?, ?, ?)
    \"\"\", (user_id, crop, disease, confidence, image_path))
    conn.commit()
    pred_id = cursor.lastrowid
    conn.close()
    return pred_id

def get_user_history(user_id: int, limit: int = 50) -> List[Dict[str, Any]]:
    conn = get_db()
    cursor = conn.cursor()
    # Fast B-Tree indexed query using idx_predictions_user_created
    cursor.execute(\"\"\"
        SELECT id, crop, disease, confidence, image_path, created_at
        FROM predictions
        WHERE user_id = ?
        ORDER BY id DESC
        LIMIT ?
    \"\"\", (user_id, limit))
    rows = cursor.fetchall()
    conn.close()
    return [dict(row) for row in rows]"""
    add_code("Snippet 5.7: Parameterized Prediction Persistence & Historical Query Functions", snippet_5_7)

    add_h2("5.5 Clinical Report Generation Engine (pdf_generator.py)")
    add_p("The clinical reporting module (pdf_generator.py) uses ReportLab's Flowable document architecture to programmatically compile official diagnostic reports. Each document contains a dual-column clinical header, a color-coded diagnostic assessment card, an embedded thumbnail of the examined leaf, a calibrated confidence meter, and comprehensive 5-pillar advisory tables:")

    snippet_5_8 = """# File: pdf_generator.py (Lines 75-115)
# Clinical PDF Construction using ReportLab Flowable Canvas
def build_clinical_pdf(prediction: dict, advisory: dict, output_stream: io.BytesIO):
    doc = SimpleDocTemplate(
        output_stream,
        pagesize=letter,
        rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36
    )
    story = []
    # Build Clinical Header with Institution & Date Metadata
    story.append(create_header_table(prediction))
    story.append(Spacer(1, 12))
    
    # Build Visual Diagnostic Assessment Card (Leaf Image + Badge)
    story.append(create_diagnosis_summary_card(prediction))
    story.append(Spacer(1, 12))
    
    # Build 5-Pillar Treatment Protocols Table
    story.append(Paragraph("<b>CLINICAL AGRONOMIC TREATMENT PRESCRIPTION</b>", header_style))
    advisory_table = Table([
        ["1. Symptoms", advisory.get("symptoms", "N/A")],
        ["2. Cause / Etiology", advisory.get("cause", "N/A")],
        ["3. Organic Remedies", advisory.get("organic_treatment", "N/A")],
        ["4. Chemical Control", advisory.get("chemical_treatment", "N/A")],
        ["5. Prevention Practices", advisory.get("prevention", "N/A")]
    ], colWidths=[120, 420])
    advisory_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (0,-1), colors.HexColor("#F1F8E9")),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#C8E6C9")),
        ('FONTNAME', (0,0), (-1,-1), 'Helvetica'),
        ('FONTSIZE', (0,0), (-1,-1), 9),
    ]))
    story.append(advisory_table)
    doc.build(story)"""
    add_code("Snippet 5.8: ReportLab Clinical PDF Document Flowable Construction", snippet_5_8)

    add_h2("5.6 Front-End User Interface & Modern Farmer Experience")
    add_p("The client-side interface (static/js/main.js) provides a reactive user experience without external JavaScript dependencies. When a foliar image is dropped into the dropzone, client-side FileReader immediately renders an image preview, activates an animated scanning pulse, dispatches the multipart request via async fetch(), and binds JSON response properties directly into the DOM:")

    snippet_5_9 = """// File: static/js/main.js (Lines 80-110)
// Reactive Client-Side Foliar Ingestion and DOM Binding
async function handleFoliarUpload(file) {
    showPreviewThumbnail(file);
    toggleLoadingState(true);
    
    const formData = new FormData();
    formData.append("file", file);
    
    try {
        const response = await fetch("/predict", { method: "POST", body: formData });
        if (!response.ok) throw new Error(await response.text());
        const data = await response.json();
        
        // Render Top-1 Primary Diagnosis and Confidence Meter
        document.getElementById("crop-name").textContent = data.crop;
        document.getElementById("disease-name").textContent = data.disease;
        updateConfidenceGauge(data.confidence);
        
        // Render Top-3 Differential Candidates List
        renderDifferentialCandidates(data.top_3_predictions);
        
        // Populate 5-Pillar Treatment Advisory Accordion
        if (data.advisory) {
            populateAdvisoryPillars(data.advisory);
        }
    } catch (err) {
        showErrorToast(err.message);
    } finally {
        toggleLoadingState(false);
    }
}"""
    add_code("Snippet 5.9: Asynchronous Fetch & Reactive DOM Binding (JavaScript)", snippet_5_9)

    doc.add_page_break()
